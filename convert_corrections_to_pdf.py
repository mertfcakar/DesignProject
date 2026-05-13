"""
Convert backend_report_corrections.md to PDF.
Uses python-docx + docx2pdf (which drives Word).

Run from project root with the project venv:
    .venv\\Scripts\\python.exe convert_corrections_to_pdf.py
"""

import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(ROOT, "backend_report_corrections.md")
DOCX_PATH = os.path.join(ROOT, "backend_report_corrections.docx")
PDF_PATH = os.path.join(ROOT, "backend_report_corrections.pdf")


def add_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def add_run(paragraph, text, bold=False, italic=False, code=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if code:
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def render_inline(paragraph, text):
    """Parse markdown inline syntax: **bold**, *italic*, `code`."""
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            add_run(paragraph, text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            add_run(paragraph, token[2:-2], bold=True)
        elif token.startswith("`") and token.endswith("`"):
            add_run(paragraph, token[1:-1], code=True, color="0B5394")
        elif token.startswith("*") and token.endswith("*"):
            add_run(paragraph, token[1:-1], italic=True)
        pos = m.end()
    if pos < len(text):
        add_run(paragraph, text[pos:])


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(6)
    if level == 1:
        run = add_run(p, text, bold=True, size=18, color="1A5C95")
        pf.page_break_before = True if "CRITICAL" not in text and "SUMMARY" not in text else False
    elif level == 2:
        add_run(p, text, bold=True, size=14, color="1A5C95")
    elif level == 3:
        add_run(p, text, bold=True, size=12, color="C75119")
    else:
        add_run(p, text, bold=True, size=11)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.paragraph_format.space_after = Pt(2)
    render_inline(p, text)


def add_numbered(doc, text, num):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, f"{num}. ", bold=True)
    render_inline(p, text)


def add_blockquote(doc, lines):
    """Render a blockquote (lines starting with >) as an indented italic block."""
    text = " ".join(line.lstrip("> ").rstrip() for line in lines if line.strip("> ").strip())
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    render_inline(p, text)
    for run in p.runs:
        run.italic = True


def add_blockquote_paragraph(doc, text):
    """Render a single paragraph of blockquote content."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    render_inline(p, text)


def add_code_block(doc, code_lines):
    """Render a code block as a single shaded paragraph in monospace."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    code_text = "\n".join(code_lines)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("1F1F1F")
    # Add light gray shading
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F4F4F4')
    p_pr.append(shd)


def add_table(doc, header_row, data_rows):
    """Render a markdown table as a real Word table."""
    n_cols = len(header_row)
    table = doc.add_table(rows=1 + len(data_rows), cols=n_cols)
    table.style = 'Light List Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header
    hdr_cells = table.rows[0].cells
    for i, txt in enumerate(header_row):
        cell = hdr_cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        add_run(p, txt, bold=True, color="FFFFFF")
        add_shading(cell, "1A5C95")

    # Data rows
    for r_idx, row in enumerate(data_rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, txt in enumerate(row):
            if c_idx >= n_cols:
                break
            cell = cells[c_idx]
            cell.text = ""
            render_inline(cell.paragraphs[0], txt.strip())


def parse_table(lines, start_idx):
    """Parse a markdown table starting at start_idx. Returns (header, data, end_idx)."""
    if start_idx + 1 >= len(lines):
        return None, None, start_idx
    header_line = lines[start_idx]
    sep_line = lines[start_idx + 1]
    if not re.match(r'^\s*\|?\s*[-:]+', sep_line):
        return None, None, start_idx

    def split_row(line):
        line = line.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        return [c.strip() for c in line.split("|")]

    header = split_row(header_line)
    data = []
    i = start_idx + 2
    while i < len(lines):
        line = lines[i]
        if not line.strip().startswith("|") and "|" not in line:
            break
        if not line.strip():
            break
        data.append(split_row(line))
        i += 1
    return header, data, i


def render_markdown_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()
    lines = text.split("\n")

    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    # Default font
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip blank lines
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^-{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("─" * 60)
            run.font.color.rgb = RGBColor.from_string("BBBBBB")
            i += 1
            continue

        # Headings
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            add_heading(doc, text, level)
            i += 1
            continue

        # Code block
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1  # skip closing ```
            continue

        # Table (line containing | and the next line being a separator)
        if "|" in stripped and i + 1 < len(lines) and re.search(r'\|.*[-:]', lines[i + 1]):
            header, data, new_i = parse_table(lines, i)
            if header:
                add_table(doc, header, data)
                i = new_i
                continue

        # Blockquote
        if stripped.startswith(">"):
            quote_text = stripped[1:].strip()
            # Detect if quote contains a markdown table or special markers
            if quote_text.startswith("**Table") or quote_text.startswith("**Listing"):
                # Render as a small caption
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                render_inline(p, quote_text)
            elif quote_text.startswith("```"):
                # Code block inside blockquote
                code_lines = []
                i += 1
                while i < len(lines):
                    inner = lines[i].lstrip(">").strip()
                    if inner.startswith("```"):
                        i += 1
                        break
                    code_lines.append(lines[i].lstrip(">").lstrip(" "))
                    i += 1
                add_code_block(doc, code_lines)
                continue
            elif "|" in quote_text and i + 1 < len(lines) and ">" in lines[i+1] and re.search(r'\|.*[-:]', lines[i+1]):
                # Try to detect a table within blockquote — strip ">" from each line
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith(">") and "|" in lines[i]:
                    table_lines.append(lines[i].lstrip(">").strip())
                    i += 1
                if len(table_lines) >= 2:
                    header, data, _ = parse_table(table_lines, 0)
                    if header:
                        add_table(doc, header, data)
                        continue
            else:
                add_blockquote_paragraph(doc, quote_text)
            i += 1
            continue

        # Bullet list
        m = re.match(r'^(\s*)[-*•]\s+(.*)$', line)
        if m:
            indent = len(m.group(1))
            level = indent // 2
            add_bullet(doc, m.group(2), level=level)
            i += 1
            continue

        # Numbered list
        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            add_numbered(doc, m.group(2), m.group(1))
            i += 1
            continue

        # Plain paragraph (collect consecutive non-blank lines)
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if nxt.startswith("#") or nxt.startswith(">") or nxt.startswith("```") or nxt.startswith("-") or nxt.startswith("*") or re.match(r'^\d+\.\s', nxt) or "|" in nxt:
                break
            paragraph_lines.append(nxt)
            i += 1

        para_text = " ".join(paragraph_lines)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        render_inline(p, para_text)

    doc.save(docx_path)
    print(f"DOCX written: {docx_path}")


def docx_to_pdf(docx_path, pdf_path):
    print("Converting DOCX to PDF via Word automation...")
    from docx2pdf import convert
    convert(docx_path, pdf_path)
    print(f"PDF written: {pdf_path}")


if __name__ == "__main__":
    print(f"Reading: {MD_PATH}")
    render_markdown_to_docx(MD_PATH, DOCX_PATH)
    docx_to_pdf(DOCX_PATH, PDF_PATH)
    print("\nDone.")
